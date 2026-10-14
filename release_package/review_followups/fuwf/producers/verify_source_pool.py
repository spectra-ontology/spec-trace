#!/usr/bin/env python3
"""Independently reread original XLSX cells; never touch production inputs."""
import argparse
import collections
import hashlib
import json
from pathlib import Path
import openpyxl


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--sampling', type=Path, required=True)
    p.add_argument('--pool', type=Path, required=True)
    p.add_argument('--output', type=Path, required=True)
    a = p.parse_args()
    sample, pool = json.loads(a.sampling.read_text()), json.loads(a.pool.read_text())
    cases = [c for c in sample['cases'] if c['relation'] == 'SENT_TO' and c['direction'] == 'source_first']
    by_path = collections.defaultdict(list)
    for case in cases:
        by_path[case['source']['source_path']].append(case)
    checked = []
    for path, entries in sorted(by_path.items()):
        before = sha(path)
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for case in entries:
            src = case['source']
            sheet = book[src['sheet']]
            headers = [str(c or '').strip() for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
            values = next(sheet.iter_rows(min_row=src['row'], max_row=src['row'], values_only=True))
            to_value = str(values[headers.index('To')] or '').strip()
            tdoc_value = str(values[headers.index('TDoc')] or '').strip()
            checked.append({'identity': case['sample_identity'], 'path': path, 'sheet': src['sheet'],
                            'row': src['row'], 'original_to': to_value, 'original_tdoc': tdoc_value,
                            'to_equal': to_value == src['to'], 'tdoc_equal': tdoc_value == src['tdoc'],
                            'sha_equal': before == src['source_read_sha256'], 'before_sha256': before})
        book.close()
        after = sha(path)
        for row in checked:
            if row['path'] == path:
                row['sha_unchanged'] = after == before
    admitted = {r['source_identity'] for r in pool['records']}
    excluded = {r['identity'] for r in pool['exclusions']}
    identities = {c['sample_identity'] for c in cases}
    valid = (len(checked) == 50 and admitted.isdisjoint(excluded) and admitted | excluded == identities
             and all(all(c[k] for k in ['to_equal', 'tdoc_equal', 'sha_equal', 'sha_unchanged']) for c in checked))
    out = {'valid': valid, 'selected_n': len(checked), 'admitted_n': len(admitted),
           'excluded_n': len(excluded), 'replacement_n': 0, 'checked': checked,
           'producer_sha256': sha(__file__), 'pool_sha256': sha(a.pool), 'sampling_sha256': sha(a.sampling),
           'expert_reviews': 0, 'model_calls': 0}
    with a.output.open('x') as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    print(json.dumps({k:v for k,v in out.items() if k != 'checked'}), flush=True)
    if not valid:
        raise ValueError('Independent original-source verification failed')


if __name__ == '__main__':
    main()
