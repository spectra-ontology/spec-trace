#!/usr/bin/env python3
"""parse_ts_docx.py — release-aware TS DOCX 파서 (RAN3, Phase-3 Stage 2)

3GPP RAN3 TS docx를 (spec, release) 전수 순회로 파싱해 Section/Table/Figure·계층·
교차참조(+TS-only Section→Spec)를 추출한다. release-aware: spec당 release별 final
전부 파싱, sectionId = {specNumber}-{release}-{sectionNumber}. release/version 정본 =
cover page(ts_cover_metadata). RAN3 고유 추출 로직(ASN.1 정의명 섹션, TH/TF 중복 ID 해소,
referencesSpec, Heading 7, "void" 부분매치)은 그대로 보존.

Spec:     docs/RAN/formal/RAN3/phase-3/design-spec/tdoc_ontology_specs(TS).md (release-aware)
Ontology: tdoc-ontology.ttl (Spec=(specNumber,release), specInRelease/specRelease)
Cycle:    release-aware-redesign (RAN1 pilot → RAN3)
배경:     기존 파서는 spec당 단일 docx(최신) 가정. release별 다버전·멀티파트로 재작성.
입력:     data/data_raw/RAN/formal/TS/RAN3/{spec}/{release}/{...}.docx (25 TS, 38.4xx 계열)
출력:     ontology/output/RAN/formal/parsed_ts/RAN3/ts_{spec}_{release}.json  (release-키, 비파괴)
          ontology/output/RAN/formal/parsed_ts/RAN3/parse_summary_release.json
사용:     python3 scripts/RAN/formal/RAN3/phase-3/ts-parser/parse_ts_docx.py [--wg RAN3]
LLM 호출: 없음
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document

PROJECT_ROOT = next(p for p in Path(__file__).resolve().parents if (p / ".git").exists())
sys.path.insert(0, str(PROJECT_ROOT / "scripts" / "RAN" / "formal" / "common" / "phase-0"))
from ts_cover_metadata import _find_doc, cover_version  # noqa: E402

RAW_DIR = PROJECT_ROOT / "data" / "data_raw" / "RAN" / "formal" / "TS"


# =============================================================================
# Configuration (release-aware: scope_specs는 main에서 디렉토리로 산출,
#                version은 cover page(ts_cover_metadata)에서 추출)
# =============================================================================


# Spec Step 5 §3.1: Heading style -> level mapping
# Spec: Heading 1~7 → level 1~7, Heading 8 → level 1 (Annex 대제목)
# RAN3 TS(38.4xx 계열): 실측 Heading 7 = 0건. 맵 엔트리는 방어적 보존이나 RAN3 실행 경로에서 매칭 0.
HEADING_LEVEL_MAP = {
    "Heading 1": 1, "Heading 2": 2, "Heading 3": 3,
    "Heading 4": 4, "Heading 5": 5, "Heading 6": 6,
    "Heading 7": 7,   # RAN3 실측 0건. 다른 WG 스타일 파일 혼입 대비 방어적 보존.
    "H6": 6,
    "Heading 8": 1,    # Annex 대제목 → level 1
    # Heading 9: RAN3 TS에 0건. 제거.
}

# Spec Step 5 §3.1: 제외 스타일
# TH(Table캡션/메시지제목), NO(NOTE), ZA/ZB/ZT/ZU(표지) — HEADING_LEVEL_MAP에 없으면 자동 스킵
ANNEX_STYLES = {"Heading 8"}  # RAN3: Heading 9는 0건

# Regex: section number validation (Spec Step 6 §13.3)
# Body 번호 규칙: 숫자(.숫자[a-zA-Z][숫자])*(.대문자)? (예: 5.3.5.13b1 패턴 허용)
BODY_NUM_RE = re.compile(r'^[0-9]+(\.[0-9]+[a-zA-Z]?[0-9]?)*(\.[A-Z])?$')
ANNEX_NUM_RE = re.compile(r'^[A-Z](\.[0-9]+)*$')
# ASN.1 정의명 패턴 — 대문자로 시작하는 CamelCase + 하이픈 결합
# RAN3 TS(38.413/423/473)에서는 ASN.1 코드가 PL 스타일로 기재되어 Section 미생성될 수 있음.
# 패턴 자체는 Spec 공통 규칙으로 보존(다른 ASN.1 정의명 헤딩 대비).
# 예: NR-RRC-Definitions, BCCH-BCH-Message, SL-MeasResultsCandRelay-r17
ASN1_NAME_RE = re.compile(r'^[A-Z][A-Za-z0-9]+([-][A-Za-z0-9]+)*$')

# Regex: Annex heading text
ANNEX_HEADING_RE = re.compile(r'^Annex\s+([A-Z])\s*(?:\([^)]*\))?\s*:?\s*(.*)')

# Regex: Table/Figure number pattern
# Handles: 5.1.3.2-1, 6.3.2.5A-1, 6.3.1.1.2-3A, A.1-1, 10.2A-1
_NUM = r'[A-Za-z0-9]+(?:\.[0-9]+[a-zA-Z]?)*(?:-[0-9]+[a-zA-Z]?)*'

# Regex: Table/Figure caption (Spec Step 7 §4.2, §4.3)
TABLE_CAPTION_RE = re.compile(rf'^Table\s+({_NUM})\s*:\s*(.+)$', re.DOTALL)
FIGURE_CAPTION_RE = re.compile(rf'^Figure\s+({_NUM})\s*:\s*(.+)$', re.DOTALL)

# Captions without colon
TABLE_NOCOLON_RE = re.compile(rf'^Table\s+({_NUM})\s*$')
TABLE_SPACE_RE = re.compile(rf'^Table\s+({_NUM})\s+(.+)$', re.DOTALL)
FIGURE_NOCOLON_RE = re.compile(rf'^Figure\s+({_NUM})\s*$')
FIGURE_SPACE_RE = re.compile(rf'^Figure\s+({_NUM})\s+(.+)$', re.DOTALL)

# Regex: Reference patterns (Spec Step 7 §4.4)
# cross-TS 패턴은 TS 38.\d{3} + TS 37.\d{3} 둘 다 허용(방어적/일반 규칙). RAN3 scope는
# 전량 38.4xx이므로 37.xxx는 scope_specs 필터에서 자연 배제되며, scope 내 TS만 참조로 해소됨.
# 패턴 1: TS + clause (Section 레벨 참조)
REF_CROSS_TS = re.compile(
    r'TS\s+(38|37)\.(\d{3})\s+(?:clause|subclause|[Ss]ection)\s+(\d+(?:\.\d+)*)')
# Spec Step 7 §4.4: TS-only 참조 (clause 없이 TS 번호만 언급) → referencesSpec (Section→Spec)
REF_CROSS_TS_ONLY = re.compile(r'TS\s+(38|37)\.(\d{3})')
REF_CLAUSE = re.compile(r'\bclause\s+(\d+(?:\.\d+)*)', re.IGNORECASE)
REF_SUBCLAUSE = re.compile(r'\bsubclause\s+(\d+(?:\.\d+)*)', re.IGNORECASE)
REF_SECTION = re.compile(r'\bSection\s+(\d+(?:\.\d+)*)')
REF_ANNEX = re.compile(r'\bAnnex\s+([A-Z](?:\.\d+)*)')


# =============================================================================
# Data classes
# =============================================================================

@dataclass
class SectionData:
    section_id: str
    section_number: str
    section_title: str
    level: int
    is_void: bool = False
    parent_section_id: str | None = None
    child_section_ids: list[str] = field(default_factory=list)
    references: list[str] = field(default_factory=list)
    references_spec: list[str] = field(default_factory=list)  # Spec §4.4: TS-only 참조 (Section→Spec)
    _body_paragraphs: list[str] = field(default_factory=list, repr=False)


@dataclass
class TableData:
    table_id: str
    table_number: str
    table_caption: str | None = None
    section_id: str | None = None


@dataclass
class FigureData:
    figure_id: str
    figure_number: str
    figure_caption: str | None = None
    section_id: str | None = None


@dataclass
class TSParseResult:
    spec_number: str
    release: str = ""                  # Rel-NN
    spec_version: str = "unknown"      # cover page 정본 (해당 release final)
    parts: list[str] = field(default_factory=list)  # 병합한 본문 docx 파트명
    sections: list[SectionData] = field(default_factory=list)
    tables: list[TableData] = field(default_factory=list)
    figures: list[FigureData] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# =============================================================================
# Parsing functions
# =============================================================================

def parse_heading_text(text: str, style: str) -> tuple[str, str]:
    """Extract sectionNumber and sectionTitle from heading text.

    Spec Step 7 §4.1:
    - Annex: "Annex A (informative): Change history" -> ("A", "Change history")
    - Body:  "5.1.3\tCSI-RS configuration" -> ("5.1.3", "CSI-RS configuration")
    """
    text = text.strip()

    # Annex heading (Heading 8/9)
    if style in ANNEX_STYLES:
        m = ANNEX_HEADING_RE.match(text)
        if m:
            return m.group(1), m.group(2).strip()
        return "", text

    # Tab-separated: "4.3\tFrame structure"
    # 확장: "–\tBCCH-BCH-Message" (em dash + tab + ASN.1 정의명)
    if '\t' in text:
        parts = text.split('\t', 1)
        sec_num = parts[0].strip()
        sec_title = parts[1].strip() if len(parts) > 1 else ""
        # em dash prefix → ASN.1 정의명이 title에 있음
        if sec_num in ('–', '–', '-'):
            # ASN.1 정의명을 sectionNumber로 사용
            return sec_title, sec_title
        # Annex "A.1:\tTitle" → 콜론 제거
        sec_num = sec_num.rstrip(':')
        return sec_num, sec_title

    # Space-separated body: "5 Physical layer procedure"
    m = re.match(r'^(\d+(?:\.\d+)*[a-zA-Z]?)\s+(.*)', text)
    if m:
        return m.group(1), m.group(2).strip()

    # Annex sub-section: "A.1 Something", "A Something", "A.1: Something"
    m = re.match(r'^([A-Z](?:\.\d+)*):?\s+(.*)', text)
    if m:
        return m.group(1), m.group(2).strip()

    # No match (e.g., "Foreword", "Introduction") -> skip
    return "", text


def _normalize(text: str) -> str:
    """Normalize paragraph text: collapse whitespace including soft line breaks."""
    return re.sub(r'\s+', ' ', text).strip()


def extract_table_info(text: str) -> tuple[str, str | None] | None:
    """Extract table number and caption from TH paragraph."""
    text = _normalize(text)
    if not text:
        return None

    m = TABLE_CAPTION_RE.match(text)
    if m:
        return m.group(1), _normalize(m.group(2))

    m = TABLE_NOCOLON_RE.match(text)
    if m:
        return m.group(1), None

    m = TABLE_SPACE_RE.match(text)
    if m:
        return m.group(1), _normalize(m.group(2))

    return None


def extract_figure_info(text: str) -> tuple[str, str | None] | None:
    """Extract figure number and caption from TF paragraph."""
    text = _normalize(text)
    if not text:
        return None

    m = FIGURE_CAPTION_RE.match(text)
    if m:
        return m.group(1), _normalize(m.group(2))

    m = FIGURE_NOCOLON_RE.match(text)
    if m:
        return m.group(1), None

    m = FIGURE_SPACE_RE.match(text)
    if m:
        return m.group(1), _normalize(m.group(2))

    return None


def number_to_section_number(num: str) -> str:
    """Map table/figure number -> containing section number.

    Spec Step 7 §4.5: remove trailing -N to get section number.
    """
    if '-' in num:
        return num.split('-', 1)[0]
    return num


def find_section_for_element(element_number: str, sid_prefix: str,
                             section_map: dict[str, SectionData]) -> str | None:
    """Find containing section for a table/figure. sid_prefix = '{spec}-{release}-'.

    Spec Step 7 §4.5:
    1. Direct match: sectionNumber = numberWithoutDash
    2. Fallback: try parent section numbers
    3. Final: return None (log warning)
    """
    target_num = number_to_section_number(element_number)
    if not target_num:
        return None

    # Direct match
    target_id = f"{sid_prefix}{target_num}"
    if target_id in section_map:
        return target_id

    # Fallback: try parent sections by removing trailing components
    parts = target_num.split('.')
    while len(parts) > 1:
        parts.pop()
        parent_num = '.'.join(parts)
        parent_id = f"{sid_prefix}{parent_num}"
        if parent_id in section_map:
            return parent_id

    # For annex sections: try just the letter
    if target_num and target_num[0].isupper() and target_num[0].isalpha():
        letter_id = f"{sid_prefix}{target_num[0]}"
        if letter_id in section_map:
            return letter_id

    return None


def build_hierarchy(sections: list[SectionData]) -> None:
    """Build parent-child section hierarchy using stack.

    Spec Step 7 §5.1: level-based parent-child relationships.
    """
    stack: list[SectionData] = []

    for section in sections:
        # Pop until we find a parent with lower level
        while stack and stack[-1].level >= section.level:
            stack.pop()

        if stack:
            parent = stack[-1]
            section.parent_section_id = parent.section_id
            parent.child_section_ids.append(section.section_id)

        stack.append(section)


def detect_void_sections(sections: list[SectionData]) -> int:
    """Detect void sections. Returns count.

    Spec Step 5 §3.1: isVoid = true when title or body only contains "Void" text.
    """
    count = 0
    for s in sections:
        # Check title
        title = s.section_title.strip().lower().rstrip('.')
        if "void" in title:
            s.is_void = True
            count += 1
            continue

        # Check body paragraphs
        non_empty = [p.strip() for p in s._body_paragraphs if p.strip()]
        if non_empty and all(p.lower().rstrip('.') == 'void' for p in non_empty):
            s.is_void = True
            count += 1

    return count


def extract_references(section: SectionData, sid_prefix: str, scope_specs: set[str],
                       all_section_ids: set[str]) -> tuple[list[str], list[str]]:
    """Extract section references and TS-only references from body text.
    동일 release 내에서만 해소 (release-aware).

    Spec Step 7 §4.4: clause/subclause/Section/Annex/cross-TS patterns.
    Returns: (references_section, references_spec)
      - references_section: Section→Section 참조 ({spec}-{release}-{num})
      - references_spec: Section→Spec TS-only 참조 (bare specNumber; 생성기가 동일 release Spec IRI로 매핑)
    """
    body_text = " ".join(section._body_paragraphs)
    if not body_text.strip():
        return [], []

    own_spec = sid_prefix.split('-Rel-')[0]
    rel_token = _release_of(section.section_id)  # 'Rel-NN' — cross-TS도 동일 release로 해소
    refs: set[str] = set()

    # Cross-TS references first (more specific)
    # cross-TS 패턴은 38.\d{3}/37.\d{3} 둘 다; scope_specs 필터로 RAN3 38.4xx만 해소됨
    for m in REF_CROSS_TS.finditer(body_text):
        target_spec = f"{m.group(1)}.{m.group(2)}"
        if target_spec in scope_specs:
            ref_id = f"{target_spec}-{rel_token}-{m.group(3)}"
            if ref_id != section.section_id and ref_id in all_section_ids:
                refs.add(ref_id)

    # Remove cross-TS matches to avoid double-matching clause refs
    cleaned = REF_CROSS_TS.sub('', body_text)

    # Same-TS references
    for pattern in [REF_CLAUSE, REF_SUBCLAUSE, REF_SECTION, REF_ANNEX]:
        for m in pattern.finditer(cleaned):
            ref_num = m.group(1)
            ref_id = f"{sid_prefix}{ref_num}"
            if ref_id != section.section_id and ref_id in all_section_ids:
                refs.add(ref_id)

    # Spec Step 7 §4.4: TS-only 참조 (clause 없이 TS 번호만) → referencesSpec
    spec_refs: set[str] = set()
    for m in REF_CROSS_TS_ONLY.finditer(body_text):
        target_spec = f"{m.group(1)}.{m.group(2)}"
        if target_spec != own_spec and target_spec in scope_specs:
            spec_refs.add(target_spec)

    return sorted(refs), sorted(spec_refs)


def _release_of(section_id: str) -> str:
    """'{spec}-{Rel-NN}-{num}' → 'Rel-NN'."""
    m = re.search(r'-(Rel-\d+)-', section_id)
    return m.group(1) if m else ""


# =============================================================================
# Main parsing logic
# =============================================================================

def find_body_parts(rel_dir: Path) -> list[Path]:
    """release 디렉토리의 본문 docx 파트. 멀티파트는 _cover 제외 전부, 단일은 그것."""
    docx = sorted(rel_dir.glob("*.docx"))
    if not docx:
        return []
    covers = [d for d in docx if d.stem.endswith("_cover")]
    if covers:
        return [d for d in docx if not d.stem.endswith("_cover")]
    s_parts = [d for d in docx if "_s" in d.stem]
    if len(docx) > 1 and s_parts:
        return s_parts
    singles = [d for d in docx if "_" not in d.stem]
    return [singles[-1]] if singles else [docx[0]]


def secnum_key(num: str) -> tuple:
    """섹션번호 자연정렬 키: 본문(0) < 부속서(1), 각 컴포넌트 (숫자, 알파)."""
    if not num:
        return (2,)
    annex = 1 if num[0].isalpha() else 0
    comps = []
    for part in num.replace('-', '.').split('.'):
        m = re.match(r'^(\d+)?([A-Za-z]*)$', part)
        if m:
            comps.append((int(m.group(1)) if m.group(1) else -1, m.group(2) or ""))
        else:
            comps.append((-1, part))
    return (annex, tuple(comps))


def parse_part_elements(docx_path: Path, sid_prefix: str, warnings: list[str],
                        seen_section_ids: set[str], table_id_set: set[str],
                        figure_id_set: set[str]) -> tuple[list[SectionData],
                                                          list[TableData], list[FigureData]]:
    """단일 docx 파트 → raw sections/tables/figures (계층/매핑 전). dedup 세트는
    (spec,release) 단위로 공유(멀티파트 대비). RAN3 고유: ASN.1 정의명 섹션 허용,
    TH/TF 중복 ID는 _suffix로 해소. sid_prefix='{spec}-{rel}-'."""
    doc = Document(str(docx_path))
    sections: list[SectionData] = []
    tables: list[TableData] = []
    figures: list[FigureData] = []
    current: SectionData | None = None

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()

        # -- Heading paragraph --
        if style_name in HEADING_LEVEL_MAP:
            if not text:
                continue

            level = HEADING_LEVEL_MAP[style_name]
            sec_num, sec_title = parse_heading_text(text, style_name)

            if not sec_num:
                continue

            # Validate section number format (ASN.1 정의명도 허용)
            valid = (BODY_NUM_RE.match(sec_num) or ANNEX_NUM_RE.match(sec_num)
                     or ASN1_NAME_RE.match(sec_num))
            if not valid:
                warnings.append(
                    f"Invalid section number: '{sec_num}' ({text[:60]})")
                continue

            sec_id = f"{sid_prefix}{sec_num}"
            if sec_id in seen_section_ids:
                warnings.append(f"Duplicate section ID: {sec_id}")
                continue

            section = SectionData(
                section_id=sec_id,
                section_number=sec_num,
                section_title=sec_title,
                level=level,
            )
            sections.append(section)
            seen_section_ids.add(sec_id)
            current = section

        # -- Table caption (TH style) --
        elif style_name == "TH":
            if not text:
                continue
            info = extract_table_info(text)
            if info:
                t_num, t_cap = info
                t_id = f"{sid_prefix}T{t_num}"
                # RAN3: 중복 tableId 방지 — suffix 추가
                if t_id in table_id_set:
                    suffix = 2
                    while f"{t_id}_{suffix}" in table_id_set:
                        suffix += 1
                    t_id = f"{t_id}_{suffix}"
                    warnings.append(
                        f"Duplicate table ID resolved: {t_id}")
                table_id_set.add(t_id)
                tables.append(TableData(
                    table_id=t_id,
                    table_number=t_num,
                    table_caption=t_cap,
                ))
            # Spec Step 5 §3.1: TH는 제외 스타일 — Table이 아닌 TH는 무시 (warning 미생성)

        # -- Figure caption (TF style) --
        elif style_name == "TF":
            if not text:
                continue
            info = extract_figure_info(text)
            if info:
                f_num, f_cap = info
                f_id = f"{sid_prefix}F{f_num}"
                # RAN3: 중복 figureId 방지 — suffix 추가
                if f_id in figure_id_set:
                    suffix = 2
                    while f"{f_id}_{suffix}" in figure_id_set:
                        suffix += 1
                    f_id = f"{f_id}_{suffix}"
                    warnings.append(
                        f"Duplicate figure ID resolved: {f_id}")
                figure_id_set.add(f_id)
                figures.append(FigureData(
                    figure_id=f_id,
                    figure_number=f_num,
                    figure_caption=f_cap,
                ))
            # TF에서 Figure 아닌 텍스트도 warning 미생성

        # -- Body paragraph -> collect for void detection + references --
        else:
            if current and text:
                current._body_paragraphs.append(text)

    return sections, tables, figures


def parse_spec_release(spec_number: str, release: str, rel_dir: Path) -> TSParseResult:
    """한 (spec, release): 본문 파트 전부 파싱 → 첫 섹션번호 정렬 병합 → 계층/매핑."""
    result = TSParseResult(spec_number=spec_number, release=release)
    sid_prefix = f"{spec_number}-{release}-"

    cover_doc = _find_doc(rel_dir)
    result.spec_version = (cover_version(cover_doc) if cover_doc else None) or "unknown"

    body_parts = find_body_parts(rel_dir)
    if not body_parts:
        result.warnings.append("No body docx parts found")
        return result

    # dedup 세트는 (spec,release) 단위로 공유 (멀티파트 간 중복 ID 해소)
    seen_section_ids: set[str] = set()
    table_id_set: set[str] = set()
    figure_id_set: set[str] = set()

    parsed_parts = []  # (sort_key, part_name, sections, tables, figures)
    for part in body_parts:
        secs, tbls, figs = parse_part_elements(
            part, sid_prefix, result.warnings,
            seen_section_ids, table_id_set, figure_id_set)
        first_num = secs[0].section_number if secs else ""
        parsed_parts.append((secnum_key(first_num), part.name, secs, tbls, figs))

    parsed_parts.sort(key=lambda x: x[0])
    for _, pname, secs, tbls, figs in parsed_parts:
        result.parts.append(pname)
        result.sections.extend(secs)
        result.tables.extend(tbls)
        result.figures.extend(figs)

    section_map = {s.section_id: s for s in result.sections}
    build_hierarchy(result.sections)
    void_count = detect_void_sections(result.sections)

    for t in result.tables:
        t.section_id = find_section_for_element(t.table_number, sid_prefix, section_map)
        if not t.section_id:
            result.warnings.append(f"Table {t.table_number}: no containing section")
    for f in result.figures:
        f.section_id = find_section_for_element(f.figure_number, sid_prefix, section_map)
        if not f.section_id:
            result.warnings.append(f"Figure {f.figure_number}: no containing section")

    print(f"  {spec_number} {release} ({len(body_parts)} part): "
          f"sec={len(result.sections)} tbl={len(result.tables)} fig={len(result.figures)} "
          f"void={void_count} V{result.spec_version}")
    return result


def to_output_dict(result: TSParseResult) -> dict:
    """Convert parse result to JSON-serializable dict (excludes body text)."""
    return {
        "spec_number": result.spec_number,
        "release": result.release,
        "specVersion": result.spec_version,
        "parts": result.parts,
        "parse_date": datetime.now().isoformat(),
        "stats": {
            "sections": len(result.sections),
            "tables": len(result.tables),
            "figures": len(result.figures),
            "void_sections": sum(1 for s in result.sections if s.is_void),
            "total_references": sum(len(s.references) for s in result.sections),
            "total_references_spec": sum(len(s.references_spec) for s in result.sections),
        },
        "sections": [
            {
                "sectionId": s.section_id,
                "sectionNumber": s.section_number,
                "sectionTitle": s.section_title,
                "level": s.level,
                "isVoid": s.is_void,
                "parentSectionId": s.parent_section_id,
                "childSectionIds": s.child_section_ids,
                "references": s.references,
                "referencesSpec": s.references_spec,
            }
            for s in result.sections
        ],
        "tables": [
            {
                "tableId": t.table_id,
                "tableNumber": t.table_number,
                "tableCaption": t.table_caption,
                "sectionId": t.section_id,
            }
            for t in result.tables
        ],
        "figures": [
            {
                "figureId": f.figure_id,
                "figureNumber": f.figure_number,
                "figureCaption": f.figure_caption,
                "sectionId": f.section_id,
            }
            for f in result.figures
        ],
        "warnings": result.warnings,
    }


# =============================================================================
# Main
# =============================================================================

def main() -> None:
    ap = argparse.ArgumentParser(description="release-aware TS DOCX 파서 (RAN3)")
    ap.add_argument("--wg", default="RAN3")
    args = ap.parse_args()
    wg = args.wg

    wg_raw = RAW_DIR / wg
    out_dir = PROJECT_ROOT / "ontology" / "output" / "RAN" / "formal" / "parsed_ts" / wg
    out_dir.mkdir(parents=True, exist_ok=True)

    print("=" * 70)
    print(f"Stage 2 (release-aware): TS DOCX Parser — {wg}")
    print("  Spec: docs/RAN/formal/RAN3/phase-3/design-spec/tdoc_ontology_specs(TS).md")
    print("=" * 70)

    rel_dirs = sorted(wg_raw.glob("*/Rel-*"))
    scope_specs = {p.name for p in wg_raw.iterdir() if p.is_dir() and p.name != "metadata"}

    # -- Pass 1: parse each (spec, release) --
    print(f"\n[Pass 1] Parsing {len(rel_dirs)} (spec,release) dirs...")
    results: list[TSParseResult] = []
    for rel_dir in rel_dirs:
        spec = rel_dir.parent.name
        release = rel_dir.name
        results.append(parse_spec_release(spec, release, rel_dir))

    # -- Pass 2: references, resolved within each release --
    print("\n[Pass 2] Extracting references (release-scoped)...")
    by_release: dict[str, list[TSParseResult]] = {}
    for r in results:
        by_release.setdefault(r.release, []).append(r)
    total_refs = 0
    total_spec_refs = 0
    for release, rlist in sorted(by_release.items()):
        ids = {s.section_id for r in rlist for s in r.sections}
        for r in rlist:
            sid_prefix = f"{r.spec_number}-{r.release}-"
            for s in r.sections:
                refs, s_refs = extract_references(s, sid_prefix, scope_specs, ids)
                s.references = refs
                s.references_spec = s_refs
            total_refs += sum(len(s.references) for s in r.sections)
            total_spec_refs += sum(len(s.references_spec) for s in r.sections)
    print(f"  Total section references: {total_refs}")
    print(f"  Total spec references (TS-only): {total_spec_refs}")

    # -- Output: one file per (spec, release) --
    print("\n[Output] Saving release-keyed parsed data...")
    for r in results:
        out_path = out_dir / f"ts_{r.spec_number}_{r.release}.json"
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(to_output_dict(r), f, ensure_ascii=False, indent=2)

    # -- Summary --
    tot_sec = sum(len(r.sections) for r in results)
    tot_tbl = sum(len(r.tables) for r in results)
    tot_fig = sum(len(r.figures) for r in results)
    tot_void = sum(1 for r in results for s in r.sections if s.is_void)
    tot_warn = sum(len(r.warnings) for r in results)
    summary = {
        "parse_date": datetime.now().isoformat(),
        "wg": wg,
        "total": {"spec_releases": len(results), "specs": len(scope_specs),
                  "sections": tot_sec, "tables": tot_tbl, "figures": tot_fig,
                  "void_sections": tot_void, "references": total_refs,
                  "references_spec": total_spec_refs, "warnings": tot_warn},
        "per_spec_release": {
            f"{r.spec_number}_{r.release}": {
                "version": r.spec_version, "parts": len(r.parts),
                "sections": len(r.sections), "tables": len(r.tables),
                "figures": len(r.figures),
                "references_spec": sum(len(s.references_spec) for s in r.sections),
                "warnings": len(r.warnings)}
            for r in results
        },
    }
    with open(out_dir / "parse_summary_release.json", "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print("\n" + "=" * 70)
    print(f"  spec-releases: {len(results)} | specs: {len(scope_specs)}")
    print(f"  Sections: {tot_sec} | Tables: {tot_tbl} | Figures: {tot_fig} | "
          f"Void: {tot_void} | Refs: {total_refs} | SpecRefs: {total_spec_refs} | "
          f"Warnings: {tot_warn}")
    print("=" * 70)


if __name__ == "__main__":
    main()
